"""
Suite de pruebas automatizadas: flujo del PRD v1.1, acceso a datos en
Supabase (con cliente simulado), exportación a Excel y endpoints.
"""
import unittest
from unittest.mock import MagicMock, patch
import os
import datetime as dt
import io

os.environ.setdefault("TELEGRAM_BOT_TOKEN", "test_token")

from app.bot_logic import (
    _remover_acentos,
    _normalizar_estado,
    _es_sin_ticket,
    generar_ticket_temporal,
    procesar_mensaje_web,
)
from app.config import ADMIN_TECNICOS
from app.state import get_estado
from app import db, excel
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from app.main import app, CATALOGOS_POR_ESTADO


class TestBotLogicPRDv11(unittest.TestCase):
    def setUp(self):
        estado_admin = get_estado("Miguel Abraham Lopez Ortiz")
        estado_admin.esperando = None
        estado_admin.folio_activo = None
        estado_admin.borrador = {}

        estado_tec = get_estado("TecnicoEstandar")
        estado_tec.esperando = None
        estado_tec.folio_activo = None
        estado_tec.borrador = {}

    def test_remover_acentos(self):
        self.assertEqual(_remover_acentos("Nivel 10"), "nivel 10")
        self.assertEqual(_remover_acentos("Términado"), "terminado")

    def test_normalizar_estado(self):
        self.assertEqual(_normalizar_estado("✅ Terminado"), "Terminado")
        self.assertEqual(_normalizar_estado("⏸️ Pendiente"), "Pendiente")
        self.assertEqual(_normalizar_estado("❌ No solucionado"), "No solucionado")

    def test_cancelar_en_cualquier_paso(self):
        estado = get_estado("TecnicoEstandar")
        estado.esperando = "ubicacion"
        estado.borrador = {"ticket": "TK-123"}

        resp = procesar_mensaje_web("TecnicoEstandar", "cancelar")
        self.assertIn("cancelada", resp[0].lower())
        self.assertIsNone(estado.esperando)
        self.assertEqual(estado.borrador, {})

    @patch("app.db.crear_reporte", return_value="#001")
    def test_flujo_completo_7_pasos(self, mock_crear):
        estado = get_estado("TecnicoEstandar")

        # 1. Inicio -> Pide proyecto / sede
        resp1 = procesar_mensaje_web("TecnicoEstandar", "➕ Nuevo reporte")
        self.assertEqual(estado.esperando, "proyecto")
        self.assertIn("paso 1", resp1[0].lower())
        self.assertIn("mina san dimas", resp1[0].lower())

        # 2. Enviar Proyecto -> Pide Ticket
        resp_p = procesar_mensaje_web("TecnicoEstandar", "Mina San Dimas")
        self.assertEqual(estado.esperando, "ticket")
        self.assertEqual(estado.borrador["proyecto"], "Mina San Dimas")
        self.assertIn("paso 2", resp_p[0].lower())

        # 3. Enviar Ticket -> Pide Departamento
        resp2 = procesar_mensaje_web("TecnicoEstandar", "TK-9988")
        self.assertEqual(estado.esperando, "departamento")
        self.assertEqual(estado.borrador["ticket"], "TK-9988")
        self.assertIn("paso 3", resp2[0].lower())

        # 4. Enviar Departamento -> Pide Ubicación / Zona
        resp3 = procesar_mensaje_web("TecnicoEstandar", "Infraestructura")
        self.assertEqual(estado.esperando, "ubicacion")
        self.assertEqual(estado.borrador["departamento"], "Infraestructura")
        self.assertIn("paso 4", resp3[0].lower())

        # 5. Enviar Ubicación -> Pide Actividad
        resp4 = procesar_mensaje_web("TecnicoEstandar", "Nivel 10")
        self.assertEqual(estado.esperando, "actividad")
        self.assertEqual(estado.borrador["ubicacion"], "Nivel 10")
        self.assertIn("paso 5", resp4[0].lower())

        # 6. Enviar Actividad -> Pide Estado
        resp5 = procesar_mensaje_web("TecnicoEstandar", "Mantenimiento a Leaky Feeder")
        self.assertEqual(estado.esperando, "estado")
        self.assertEqual(estado.borrador["actividad"], "Mantenimiento a Leaky Feeder")
        self.assertIn("paso 6", resp5[0].lower())

        # 7. Enviar Estado -> Pide Evidencia
        resp6 = procesar_mensaje_web("TecnicoEstandar", "Terminado")
        self.assertEqual(estado.esperando, "evidencia")
        self.assertEqual(estado.borrador["estado"], "Terminado")
        self.assertIn("paso 7", resp6[0].lower())

        # 8. Omitir Evidencia -> Pide Confirmación
        resp7 = procesar_mensaje_web("TecnicoEstandar", "Omitir")
        self.assertEqual(estado.esperando, "confirmacion")
        self.assertIn("confirmación", resp7[0].lower())
        self.assertIn("Infraestructura", resp7[0])
        self.assertIn("Mina San Dimas", resp7[0])

        # 9. Guardar -> Llama a db.crear_reporte y confirma con #001
        resp8 = procesar_mensaje_web("TecnicoEstandar", "Guardar")
        self.assertIsNone(estado.esperando)
        self.assertEqual(estado.borrador, {})
        mock_crear.assert_called_once_with(
            ticket="TK-9988",
            tecnico="TecnicoEstandar",
            ubicacion="Nivel 10",
            actividad="Mantenimiento a Leaky Feeder",
            estado="Terminado",
            departamento="Infraestructura",
            evidencias=[],
            proyecto="Mina San Dimas",
        )
        self.assertIn("#001", resp8[0])
        self.assertIn("• *Departamento:* Infraestructura", resp8[0])
        self.assertIn("• *Proyecto:* Mina San Dimas", resp8[0])

    def test_generar_ticket_temporal(self):
        t = generar_ticket_temporal()
        self.assertTrue(t.startswith("S/T-"))
        self.assertGreaterEqual(len(t), 10)

    def test_es_sin_ticket(self):
        self.assertTrue(_es_sin_ticket("Sin ticket"))
        self.assertTrue(_es_sin_ticket("sin ticket"))
        self.assertTrue(_es_sin_ticket("S/T"))
        self.assertTrue(_es_sin_ticket("s/t"))
        self.assertTrue(_es_sin_ticket("st"))
        self.assertTrue(_es_sin_ticket("No tengo ticket"))
        self.assertTrue(_es_sin_ticket("omitir"))
        self.assertTrue(_es_sin_ticket("N/A"))
        self.assertTrue(_es_sin_ticket("ninguno"))
        self.assertFalse(_es_sin_ticket("TK-001"))
        self.assertFalse(_es_sin_ticket("12345"))

    def test_catalogos_por_estado_ticket(self):
        self.assertIn("Sin ticket", CATALOGOS_POR_ESTADO.get("ticket", []))

    @patch("app.db.crear_reporte", return_value="#002")
    def test_flujo_sin_ticket_guardado_completo(self, mock_crear):
        estado = get_estado("TecnicoEstandar")

        # 1. Inicio -> Pide proyecto / sede
        resp1 = procesar_mensaje_web("TecnicoEstandar", "➕ Nuevo reporte")
        self.assertEqual(estado.esperando, "proyecto")
        self.assertIn("paso 1", resp1[0].lower())

        # 2. Proyecto -> Pide ticket (con opción Sin ticket)
        resp_p = procesar_mensaje_web("TecnicoEstandar", "Proyectos Durango")
        self.assertEqual(estado.esperando, "ticket")
        self.assertIn("sin ticket", resp_p[0].lower())

        # 3. Responde 'Sin ticket' -> asigna ticket temporal y pide Departamento
        resp2 = procesar_mensaje_web("TecnicoEstandar", "Sin ticket")
        self.assertEqual(estado.esperando, "departamento")
        ticket_asignado = estado.borrador["ticket"]
        self.assertTrue(ticket_asignado.startswith("S/T-"))
        self.assertIn("ticket temporal", resp2[0].lower())
        self.assertIn("paso 3", resp2[1].lower())

        # 4. Departamento -> Ubicación
        resp3 = procesar_mensaje_web("TecnicoEstandar", "Infraestructura")
        self.assertEqual(estado.esperando, "ubicacion")
        self.assertIn("paso 4", resp3[0].lower())

        # 5. Ubicación -> Actividad
        procesar_mensaje_web("TecnicoEstandar", "Oficinas Centrales")
        self.assertEqual(estado.esperando, "actividad")

        # 6. Actividad -> Estado
        procesar_mensaje_web("TecnicoEstandar", "Revisión de radio")
        self.assertEqual(estado.esperando, "estado")

        # 7. Estado -> Evidencia
        procesar_mensaje_web("TecnicoEstandar", "Terminado")
        self.assertEqual(estado.esperando, "evidencia")

        # 8. Omitir Evidencia -> Confirmación
        resp_conf = procesar_mensaje_web("TecnicoEstandar", "Omitir")
        self.assertEqual(estado.esperando, "confirmacion")
        self.assertIn(ticket_asignado, resp_conf[0])
        self.assertIn("Infraestructura", resp_conf[0])
        self.assertIn("Proyectos Durango", resp_conf[0])

        # 9. Guardar -> db.crear_reporte con el ticket temporal generado y proyecto
        resp_guardar = procesar_mensaje_web("TecnicoEstandar", "Guardar")
        self.assertIsNone(estado.esperando)
        self.assertEqual(estado.borrador, {})
        mock_crear.assert_called_once_with(
            ticket=ticket_asignado,
            tecnico="TecnicoEstandar",
            ubicacion="Oficinas Centrales",
            actividad="Revisión de radio",
            estado="Terminado",
            departamento="Infraestructura",
            evidencias=[],
            proyecto="Proyectos Durango",
        )
        self.assertIn("#002", resp_guardar[0])

    @patch("app.db.agregar_tecnico", return_value="ABC123")
    def test_admin_flujo_nuevo_tecnico(self, mock_agregar):
        estado = get_estado("Miguel Abraham Lopez Ortiz")

        resp = procesar_mensaje_web("Miguel Abraham Lopez Ortiz", "dar de alta")
        self.assertEqual(estado.esperando, "admin_nombre_tecnico")

        resp2 = procesar_mensaje_web("Miguel Abraham Lopez Ortiz", "Juan Pérez")
        self.assertIsNone(estado.esperando)
        mock_agregar.assert_called_once_with("Juan Pérez")
        self.assertIn("/start abc123", resp2[0].lower())

    def test_no_admin_bloqueado(self):
        resp = procesar_mensaje_web("TecnicoEstandar", "nuevo tecnico")
        self.assertIn("no tienes permisos", resp[0].lower())

    @patch("app.db.listar_reportes_pendientes")
    @patch("app.db.actualizar_reporte_pendiente", return_value=True)
    def test_flujo_continuar_pendiente(self, mock_actualizar, mock_pendientes):
        mock_pendientes.return_value = [
            {
                "numero": "#002",
                "ticket": "TK-500",
                "ubicacion": "Nivel 11",
                "actividad": "Instalación de switch",
                "estado": "Pendiente",
                "fecha": "2026-09-15",
            }
        ]

        estado = get_estado("TecnicoEstandar")

        # 1. Consultar mis pendientes
        resp1 = procesar_mensaje_web("TecnicoEstandar", "Mis pendientes")
        self.assertEqual(estado.esperando, "seleccion_pendiente")
        self.assertIn("#002", resp1[0])

        # 2. Seleccionar reporte #002 -> Pide descripción
        resp2 = procesar_mensaje_web("TecnicoEstandar", "002")
        self.assertEqual(estado.esperando, "continuar_pendiente_actividad")
        self.assertEqual(estado.folio_activo, "#002")

        # 3. Enviar descripción -> Pide nuevo estado
        resp3 = procesar_mensaje_web("TecnicoEstandar", "Se configuró VLAN y quedó probado")
        self.assertEqual(estado.esperando, "continuar_pendiente_estado")

        # 4. Elegir estado Terminado -> Actualiza en la base de datos
        resp4 = procesar_mensaje_web("TecnicoEstandar", "Terminado")
        self.assertIsNone(estado.esperando)
        mock_actualizar.assert_called_once_with(
            "#002",
            tecnico="TecnicoEstandar",
            nueva_actividad="Se configuró VLAN y quedó probado",
            nuevo_estado="Terminado",
        )
        self.assertIn("actualizado", resp4[0].lower())

    @patch("app.db.listar_reportes_pendientes")
    def test_no_puede_elegir_reporte_fuera_de_sus_pendientes(self, mock_pendientes):
        mock_pendientes.return_value = [{
            "numero": "#002", "ticket": "TK-500", "ubicacion": "Nivel 11",
            "actividad": "x", "estado": "Pendiente", "fecha": "2026-09-15",
        }]
        estado = get_estado("TecnicoEstandar")
        procesar_mensaje_web("TecnicoEstandar", "Mis pendientes")
        resp = procesar_mensaje_web("TecnicoEstandar", "#007")
        self.assertEqual(estado.esperando, "seleccion_pendiente")
        self.assertIn("no está en tu lista", resp[0])

    def test_estado_invalido_vuelve_a_preguntar(self):
        estado = get_estado("TecnicoEstandar")
        estado.esperando = "estado"
        estado.borrador = {"evidencias": []}
        resp = procesar_mensaje_web("TecnicoEstandar", "más o menos")
        self.assertEqual(estado.esperando, "estado")
        self.assertIn("Elige uno de estos estados", resp[0])

    @patch("app.db.asegurar_zona", return_value="Nivel 15")
    def test_zona_otra_auto_registra(self, mock_asegurar):
        estado = get_estado("TecnicoEstandar")
        estado.esperando = "ubicacion"
        estado.borrador = {"ticket": "TK-1", "departamento": "Infraestructura"}

        resp = procesar_mensaje_web("TecnicoEstandar", "Otra")
        self.assertEqual(estado.esperando, "ubicacion_otra")
        self.assertIn("personalizada", resp[0].lower())

        resp2 = procesar_mensaje_web("TecnicoEstandar", "Nivel 15")
        self.assertEqual(estado.esperando, "actividad")
        self.assertEqual(estado.borrador["ubicacion"], "Nivel 15")
        mock_asegurar.assert_called_once_with("Nivel 15")

    def test_departamento_invalido_repregunta(self):
        estado = get_estado("TecnicoEstandar")
        estado.esperando = "departamento"
        estado.borrador = {"ticket": "TK-1"}

        with patch("app.db.listar_departamentos", return_value=["Infraestructura"]):
            resp = procesar_mensaje_web("TecnicoEstandar", "Inexistente")
            self.assertEqual(estado.esperando, "departamento")
            self.assertIn("válido", resp[0].lower())


class TestBotonesConEmoji(unittest.TestCase):
    """Los botones de Telegram/Web llegan con emoji (algunos con el selector
    de variación U+FE0F, ej. "✏️"); bot_logic debe reconocerlos igual."""

    def setUp(self):
        for nombre in ("Miguel Abraham Lopez Ortiz", "TecnicoEstandar"):
            estado = get_estado(nombre)
            estado.esperando = None
            estado.folio_activo = None
            estado.borrador = {}

    def test_remover_acentos_quita_emoji(self):
        self.assertEqual(_remover_acentos("✏️ Editar"), "editar")
        self.assertEqual(_remover_acentos("⏸️ Mis pendientes"), "mis pendientes")
        self.assertEqual(_remover_acentos("👤 + Nuevo técnico"), "+ nuevo tecnico")

    @patch("app.db.listar_mis_reportes", return_value=[])
    @patch("app.db.listar_reportes_pendientes", return_value=[])
    def test_todos_los_botones_del_teclado_se_reconocen(self, _pend, _rep):
        from app.telegram_client import TECLADO_ADMIN
        botones = [b["text"] for fila in TECLADO_ADMIN["keyboard"] for b in fila
                   if b["text"] != "📊 Exportar Excel"]
        for boton in botones:
            get_estado("Miguel Abraham Lopez Ortiz").esperando = None
            resp = procesar_mensaje_web("Miguel Abraham Lopez Ortiz", boton)
            self.assertNotIn("no entendí", resp[0].lower(), boton)

    def test_botones_teclado_anterior_redirigen(self):
        from app.telegram_client import normalizar_texto_boton
        self.assertEqual(normalizar_texto_boton("+ Nueva actividad"), "nuevo reporte")
        self.assertEqual(normalizar_texto_boton("✓ Finalizar"), "mis pendientes")

    def test_confirmacion_boton_editar(self):
        estado = get_estado("TecnicoEstandar")
        estado.esperando = "confirmacion"
        estado.borrador = {"ticket": "TK-1", "evidencias": []}
        resp = procesar_mensaje_web("TecnicoEstandar", "✏️ Editar")
        self.assertEqual(estado.esperando, "editar_campo_seleccion")
        self.assertIn("corregir", resp[0].lower())

        # Probar selección de ticket y modificación
        resp_sel = procesar_mensaje_web("TecnicoEstandar", "Ticket")
        self.assertEqual(estado.esperando, "editar_ticket")

        resp_mod = procesar_mensaje_web("TecnicoEstandar", "TK-NUEVO-99")
        self.assertEqual(estado.esperando, "confirmacion")
        self.assertEqual(estado.borrador["ticket"], "TK-NUEVO-99")

    def test_ubicacion_solo_emoji_no_elige_nivel_10(self):
        estado = get_estado("TecnicoEstandar")
        estado.esperando = "ubicacion"
        estado.borrador = {"evidencias": []}
        procesar_mensaje_web("TecnicoEstandar", "🏗️")
        self.assertEqual(estado.borrador["ubicacion"], "🏗️")


def _fila_db(id_=1, estado="Terminado", evidencias=None, tecnico="Juan"):
    return {
        "id": id_, "ticket": f"TK-{id_}", "tecnico": tecnico, "ubicacion": "Nivel 10",
        "actividad": "Cambio de AP", "estado": estado, "evidencias": evidencias or [],
        "creado": "2026-09-22T17:05:09.123456+00:00",
        "actualizado": "2026-09-22T17:05:09.123456+00:00",
    }


class TestDB(unittest.TestCase):
    def test_numero_formato_y_parseo(self):
        self.assertEqual(db.formatear_numero(3), "#003")
        self.assertEqual(db.formatear_numero(1234), "#1234")
        self.assertEqual(db.parsear_numero("#002"), 2)
        self.assertEqual(db.parsear_numero(" 15 "), 15)
        self.assertIsNone(db.parsear_numero("abc"))

    def test_fila_a_reporte_en_hora_local(self):
        r = db._reporte_desde_fila(_fila_db(id_=7, evidencias=["https://x/1.jpg"]))
        self.assertEqual(r["numero"], "#007")
        # 17:05 UTC = 11:05 en America/Mexico_City (UTC-6)
        self.assertEqual((r["fecha"], r["hora"]), ("2026-09-22", "11:05:09"))
        self.assertEqual(r["evidencias"], ["https://x/1.jpg"])

    @patch("app.db.get_client")
    def test_crear_reporte_regresa_numero_del_id(self, mock_client):
        tabla = mock_client.return_value.table.return_value
        tabla.insert.return_value.execute.return_value.data = [{"id": 42}]
        numero = db.crear_reporte("TK-1", "Juan", "Nivel 10", "x", "Terminado", ["u"])
        self.assertEqual(numero, "#042")
        self.assertEqual(tabla.insert.call_args[0][0]["evidencias"], ["u"])

    @patch("app.db.get_client")
    def test_actualizar_pendiente_ajeno_no_escribe(self, mock_client):
        tabla = mock_client.return_value.table.return_value
        consulta = tabla.select.return_value.eq.return_value.eq.return_value.eq.return_value.limit.return_value
        consulta.execute.return_value.data = []
        ok = db.actualizar_reporte_pendiente("#003", "Otro", "avance", "Terminado")
        self.assertFalse(ok)
        tabla.update.assert_not_called()

    @patch("app.db.get_client")
    def test_reportes_en_periodo_usa_dias_locales(self, mock_client):
        tabla = mock_client.return_value.table.return_value
        consulta = tabla.select.return_value.gte.return_value.lt.return_value.order.return_value
        consulta.execute.return_value.data = [_fila_db()]
        reportes = db.reportes_en_periodo(dt.date(2026, 9, 21), dt.date(2026, 9, 27))
        self.assertEqual(len(reportes), 1)
        tabla.select.return_value.gte.assert_called_once_with("creado", "2026-09-21T00:00:00-06:00")
        tabla.select.return_value.gte.return_value.lt.assert_called_once_with("creado", "2026-09-28T00:00:00-06:00")


def _png(ancho=400, alto=300) -> bytes:
    from PIL import Image as PILImage
    buf = io.BytesIO()
    PILImage.new("RGB", (ancho, alto), "red").save(buf, format="PNG")
    return buf.getvalue()


class TestExcel(unittest.TestCase):
    @patch("app.excel._descargar", side_effect=lambda url: _png())
    def test_genera_hojas_y_datos(self, _descargar):
        reportes = [
            db._reporte_desde_fila(_fila_db(1, "Terminado", ["https://x/1.jpg", "https://x/2.jpg"])),
            db._reporte_desde_fila(_fila_db(2, "Pendiente", tecnico="Ana")),
        ]
        contenido, nombre = excel.generar_excel(reportes, dt.date(2026, 9, 21), dt.date(2026, 9, 27))
        self.assertEqual(nombre, "Reporte_Campo_2026-09-21_a_2026-09-27.xlsx")

        wb = load_workbook(io.BytesIO(contenido))
        self.assertEqual(wb.sheetnames, ["Resumen", "Reportes", "Evidencias"])
        ws = wb["Reportes"]
        self.assertEqual(ws["A2"].value, "#001")
        self.assertEqual(ws["C1"].value, "Departamento")
        self.assertEqual(ws["C2"].value, "Infraestructura")
        self.assertEqual(ws["F3"].value, "Pendiente")
        self.assertEqual(ws["G2"].value, 2)  # número de fotos
        self.assertEqual(ws["K1"].value, "Evidencias")
        # Dos miniaturas ancladas a la celda K2 del reporte #001
        self.assertEqual(len(ws._images), 2)
        self.assertTrue(all(img.anchor._from.row == 1 and img.anchor._from.col == 10 for img in ws._images))
        self.assertTrue(all(img.height == 120 for img in ws._images))
        ev = wb["Evidencias"]
        self.assertEqual(ev["D3"].value, "https://x/2.jpg")
        self.assertEqual(ev["D3"].hyperlink.target, "https://x/2.jpg")
        resumen = [c.value for fila in wb["Resumen"].iter_rows() for c in fila]
        self.assertIn("Total de reportes", resumen)
        self.assertIn("Por departamento", resumen)
        self.assertIn("Infraestructura", resumen)

    @patch("app.excel._descargar", side_effect=[b"%PDF-1.4 no es imagen", None])
    def test_evidencia_no_imagen_o_caida(self, _descargar):
        reportes = [db._reporte_desde_fila(_fila_db(1, evidencias=["https://x/a.pdf", "https://x/b.jpg"]))]
        contenido, _ = excel.generar_excel(reportes, dt.date(2026, 9, 21), dt.date(2026, 9, 27))
        ws = load_workbook(io.BytesIO(contenido))["Reportes"]
        self.assertEqual(len(ws._images), 0)
        self.assertEqual(ws["K2"].value, "Ver hoja Evidencias")

    @patch("app.excel.httpx.get")
    @patch("app.storage.get_client")
    def test_descarga_por_storage_aunque_bucket_privado(self, mock_client, mock_get):
        bucket = mock_client.return_value.storage.from_.return_value
        bucket.download.return_value = b"jpg"
        url = "https://abc.supabase.co/storage/v1/object/public/evidencias/evidencia_Juan_file%201.jpg"
        self.assertEqual(excel._descargar(url), b"jpg")
        bucket.download.assert_called_once_with("evidencia_Juan_file 1.jpg")
        mock_get.assert_not_called()

    @patch("app.excel.httpx.get")
    @patch("app.storage.get_client", side_effect=RuntimeError("sin supabase"))
    def test_descarga_cae_a_link_publico(self, _client, mock_get):
        mock_get.return_value.content = b"jpg"
        self.assertEqual(excel._descargar("https://otro.sitio/foto.jpg"), b"jpg")

    @patch("app.excel._descargar", side_effect=lambda url: _png())
    def test_misma_foto_en_varios_reportes(self, _descargar):
        reportes = [db._reporte_desde_fila(_fila_db(i, evidencias=["https://x/a.jpg"])) for i in (1, 2)]
        contenido, _ = excel.generar_excel(reportes, dt.date(2026, 9, 21), dt.date(2026, 9, 27))
        self.assertEqual(len(load_workbook(io.BytesIO(contenido))["Reportes"]._images), 2)

    def test_periodo_vacio(self):
        contenido, _ = excel.generar_excel([], dt.date(2026, 1, 1), dt.date(2026, 1, 7))
        wb = load_workbook(io.BytesIO(contenido))
        self.assertEqual(wb["Reportes"].max_row, 1)

    def test_excel_multiples_departamentos(self):
        reportes = [
            {"numero": "#001", "ticket": "TK-1", "tecnico": "Juan", "departamento": "Infraestructura",
             "ubicacion": "Nivel 10", "actividad": "Act 1", "estado": "Terminado", "evidencias": [],
             "fecha": "2026-09-22", "hora": "10:00:00", "actualizado": "2026-09-22 10:00:00"},
            {"numero": "#002", "ticket": "TK-2", "tecnico": "Ana", "departamento": "Mantenimiento",
             "ubicacion": "Nivel 11", "actividad": "Act 2", "estado": "Pendiente", "evidencias": [],
             "fecha": "2026-09-22", "hora": "11:00:00", "actualizado": "2026-09-22 11:00:00"},
        ]
        contenido, _ = excel.generar_excel(reportes, dt.date(2026, 9, 21), dt.date(2026, 9, 27))
        wb = load_workbook(io.BytesIO(contenido))
        ws = wb["Reportes"]
        self.assertEqual(ws["C2"].value, "Infraestructura")
        self.assertEqual(ws["C3"].value, "Mantenimiento")
        resumen = [c.value for fila in wb["Resumen"].iter_rows() for c in fila]
        self.assertIn("Mantenimiento", resumen)


class TestFastAPIEndpoints(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def test_health_check(self):
        res = self.client.get("/health")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json(), {"status": "healthy"})

    @patch("app.main.WEBHOOK_SECRET", "test_secret")
    @patch("app.main._tecnico_de", return_value="TecnicoEstandar")
    @patch("app.telegram_client.send_opciones")
    @patch("app.telegram_client.send_text")
    def test_telegram_webhook(self, mock_send, mock_opciones, mock_tecnico):
        res = self.client.post(
            "/telegram-webhook",
            json={"message": {"chat": {"id": 12345}, "text": "Hola"}},
            headers={"X-Telegram-Bot-Api-Secret-Token": "test_secret"},
        )
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json(), {"status": "ok"})

    @patch("app.main.WEBHOOK_SECRET", "test_secret")
    @patch("app.main._tecnico_de", return_value="TecnicoEstandar")
    @patch("app.telegram_client.descargar_archivo", return_value=(_png(), "evidencias/test.png"))
    @patch("app.storage.upload_evidence", return_value="https://storage/test.png")
    @patch("app.telegram_client.send_opciones")
    @patch("app.telegram_client.send_text")
    def test_telegram_webhook_foto(self, mock_send_text, mock_send_opciones, mock_upload, mock_download, mock_tecnico):
        estado = get_estado("TecnicoEstandar")
        estado.esperando = "evidencia"
        estado.borrador = {"evidencias": []}

        res = self.client.post(
            "/telegram-webhook",
            json={"message": {"chat": {"id": 12345}, "photo": [{"file_id": "file_123"}]}},
            headers={"X-Telegram-Bot-Api-Secret-Token": "test_secret"},
        )
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json(), {"status": "ok"})
        mock_send_opciones.assert_called_once()



class TestSeguridadEndpoints(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    @patch("app.main.REPORTE_ADMIN_SECRET", None)
    def test_admin_sin_secreto_configurado_rechaza(self):
        # Falla en cerrado: sin REPORTE_ADMIN_SECRET nadie entra, ni con ?secret= vacío
        for url in ["/api/exportar-excel", "/api/codigo-activacion?nombre=X"]:
            res = self.client.get(url)
            self.assertEqual(res.status_code, 403, url)

    @patch("app.main.REPORTE_ADMIN_SECRET", "s3creto")
    def test_admin_secreto_incorrecto_rechaza(self):
        res = self.client.get("/api/codigo-activacion?secret=otro&nombre=X")
        self.assertEqual(res.status_code, 403)

    @patch("app.main.REPORTE_ADMIN_SECRET", "s3creto")
    @patch("app.db.codigo_activacion_pendiente", return_value="ABC123")
    def test_admin_secreto_correcto_query_o_header(self, _mock):
        res = self.client.get("/api/codigo-activacion?secret=s3creto&nombre=X")
        self.assertEqual(res.status_code, 200)
        res = self.client.get("/api/codigo-activacion?nombre=X", headers={"X-Admin-Secret": "s3creto"})
        self.assertEqual(res.status_code, 200)

    @patch("app.main.REPORTE_ADMIN_SECRET", "s3creto")
    @patch("app.db.reportes_en_periodo", return_value=[])
    def test_exportar_excel_con_secreto(self, mock_periodo):
        res = self.client.get("/api/exportar-excel?desde=2026-09-01&hasta=2026-09-30",
                              headers={"X-Admin-Secret": "s3creto"})
        self.assertEqual(res.status_code, 200)
        self.assertIn("Reporte_Campo_2026-09-01_a_2026-09-30.xlsx", res.headers["content-disposition"])
        mock_periodo.assert_called_once_with(dt.date(2026, 9, 1), dt.date(2026, 9, 30))

    @patch("app.main.REPORTE_ADMIN_SECRET", "s3creto")
    def test_exportar_excel_fechas_invalidas(self):
        res = self.client.get("/api/exportar-excel?secret=s3creto&desde=2026-09-30&hasta=2026-09-01")
        self.assertEqual(res.status_code, 400)

    @patch("app.main.WEBAPP_HABILITADA", False)
    def test_webapp_deshabilitada_por_default(self):
        res = self.client.post("/api/chat", json={"tecnico": "Miguel Abraham Lopez Ortiz", "texto": "nuevo tecnico"})
        self.assertEqual(res.status_code, 404)
        self.assertEqual(self.client.get("/api/tecnicos").status_code, 404)
        self.assertEqual(self.client.get("/").json()["status"], "ok")

    @patch("app.main.WEBHOOK_SECRET", "hook")
    def test_webhook_secreto_incorrecto(self):
        res = self.client.post(
            "/telegram-webhook",
            json={"message": {"chat": {"id": 1}, "text": "hola"}},
            headers={"X-Telegram-Bot-Api-Secret-Token": "otro"},
        )
        self.assertEqual(res.status_code, 403)



class TestComandoReporteTelegram(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def _mandar(self, texto):
        with patch("app.main.WEBHOOK_SECRET", "test_secret"):
            return self.client.post(
                "/telegram-webhook",
                json={"message": {"chat": {"id": 1}, "text": texto}},
                headers={"X-Telegram-Bot-Api-Secret-Token": "test_secret"},
            )

    @patch("app.main._tecnico_de", return_value="Miguel Abraham Lopez Ortiz")
    @patch("app.db.reportes_en_periodo", return_value=[])
    @patch("app.telegram_client.send_document")
    @patch("app.telegram_client.send_text")
    def test_admin_recibe_excel(self, _texto, mock_doc, mock_periodo, _tec):
        self._mandar("/reporte 2026-09-01 2026-09-15")
        mock_periodo.assert_called_once_with(dt.date(2026, 9, 1), dt.date(2026, 9, 15))
        args, kwargs = mock_doc.call_args
        self.assertTrue(args[2].endswith(".xlsx"))
        self.assertIn("spreadsheetml", kwargs["mime_type"])

    @patch("app.main._tecnico_de", return_value="Miguel Abraham Lopez Ortiz")
    @patch("app.db.reportes_en_periodo", return_value=[])
    @patch("app.telegram_client.send_document")
    @patch("app.telegram_client.send_text")
    def test_boton_exportar_excel(self, _texto, mock_doc, _periodo, _tec):
        self._mandar("📊 Exportar Excel")
        mock_doc.assert_called_once()

    @patch("app.main._tecnico_de", return_value="TecnicoEstandar")
    @patch("app.telegram_client.send_document")
    @patch("app.telegram_client.send_text")
    def test_no_admin_no_recibe_excel(self, mock_texto, mock_doc, _tec):
        self._mandar("/reporte")
        mock_doc.assert_not_called()
        self.assertIn("No tienes permiso", mock_texto.call_args[0][1])

    @patch("app.main._tecnico_de", return_value="Miguel Abraham Lopez Ortiz")
    @patch("app.db.listar_mis_reportes", return_value=[])
    @patch("app.telegram_client.send_document")
    @patch("app.telegram_client.send_text")
    def test_mis_reportes_no_es_el_comando_de_admin(self, _texto, mock_doc, mock_mis, _tec):
        self._mandar("/reportes")
        mock_doc.assert_not_called()
        mock_mis.assert_called_once()

    @patch("app.main._tecnico_de", return_value="Miguel Abraham Lopez Ortiz")
    @patch("app.db.agregar_departamento", return_value=True)
    @patch("app.telegram_client.send_text")
    def test_admin_crea_departamento(self, mock_texto, mock_db, _tec):
        self._mandar("/nuevo_departamento Mantenimiento")
        mock_db.assert_called_once_with("Mantenimiento")
        self.assertIn("creado", mock_texto.call_args[0][1].lower())

    @patch("app.main._tecnico_de", return_value="Miguel Abraham Lopez Ortiz")
    @patch("app.db.listar_departamentos", return_value=["Infraestructura", "Mantenimiento"])
    @patch("app.telegram_client.send_opciones")
    def test_admin_lista_departamentos(self, mock_opciones, mock_db, _tec):
        self._mandar("/departamentos")
        mock_db.assert_called_once()
        self.assertIn("Mantenimiento", mock_opciones.call_args[0][1])

    @patch("app.main._tecnico_de", return_value="Miguel Abraham Lopez Ortiz")
    @patch("app.db.listar_proyectos", return_value=[
        {"nombre": "Mina San Dimas", "tipo": "Mina", "cliente": "First Majestic"},
        {"nombre": "Proyectos Durango", "tipo": "Durango", "cliente": "Qtek Computación"},
    ])
    @patch("app.telegram_client.send_opciones")
    def test_admin_lista_proyectos_boton_teclado(self, mock_opciones, mock_db, _tec):
        self._mandar("🏗️ Proyectos")
        mock_db.assert_called_once()
        self.assertIn("Mina San Dimas", mock_opciones.call_args[0][1])
        self.assertIn("Proyectos Durango", mock_opciones.call_args[0][1])

    def test_one_shot_registro_rapido(self):
        estado = get_estado("TecnicoEstandar")
        with patch("app.db.asegurar_zona", return_value="Nivel 10"):
            resp = procesar_mensaje_web(
                "TecnicoEstandar",
                "/registro TK-999 | Infraestructura | Nivel 10 | Cambio de patchcord | Terminado",
            )
        self.assertEqual(estado.esperando, "confirmacion")
        self.assertEqual(estado.borrador["ticket"], "TK-999")
        self.assertEqual(estado.borrador["departamento"], "Infraestructura")
        self.assertEqual(estado.borrador["ubicacion"], "Nivel 10")
        self.assertEqual(estado.borrador["actividad"], "Cambio de patchcord")
        self.assertEqual(estado.borrador["estado"], "Terminado")
        self.assertIn("confirmación", resp[0].lower())

    @patch("app.db.buscar_reporte")
    def test_comando_buscar_reporte(self, mock_buscar):
        mock_buscar.return_value = [{
            "numero": "#001",
            "ticket": "TK-100",
            "departamento": "Infraestructura",
            "ubicacion": "Nivel 10",
            "actividad": "Reparación",
            "estado": "Terminado",
            "tecnico": "Miguel",
            "fecha": "2026-10-14",
            "evidencias": [],
        }]
        resp = procesar_mensaje_web("TecnicoEstandar", "/buscar TK-100")
        mock_buscar.assert_called_once_with("TK-100")
        self.assertIn("TK-100", resp[0])
        self.assertIn("#001", resp[0])

    @patch("app.db.listar_reportes_por_zona")
    def test_comando_zona_reportes(self, mock_zona):
        mock_zona.return_value = [{
            "numero": "#002",
            "ticket": "TK-200",
            "actividad": "Revisión Leaky Feeder",
            "estado": "Pendiente",
            "fecha": "2026-10-14",
        }]
        resp = procesar_mensaje_web("TecnicoEstandar", "/zona Nivel 10")
        mock_zona.assert_called_once_with("Nivel 10")
        self.assertIn("Nivel 10", resp[0])
        self.assertIn("#002", resp[0])

    def test_marca_de_agua_procesamiento(self):
        from app.watermark import aplicar_marca_de_agua
        from PIL import Image
        img = Image.new("RGB", (300, 200), color="red")
        b = io.BytesIO()
        img.save(b, "JPEG")
        resultado = aplicar_marca_de_agua(b.getvalue(), ticket="TK-123", tecnico="Juan", zona="Nivel 11")
        self.assertIsInstance(resultado, bytes)
        self.assertGreater(len(resultado), 0)

    def test_generar_zip_evidencias(self):
        from app.zip_export import generar_zip_evidencias
        zip_bytes, nombre = generar_zip_evidencias([], dt.date.today(), dt.date.today())
        self.assertTrue(nombre.endswith(".zip"))
        self.assertGreater(len(zip_bytes), 0)

    @patch("app.db.crear_reporte", return_value="#099")
    def test_registro_rapido_con_proyecto_durango(self, mock_crear):
        resp = procesar_mensaje_web(
            "TecnicoEstandar",
            "/registro Proyectos Durango | TK-500 | Sistemas | Oficinas Centrales | Instalación Switch | Terminado"
        )
        self.assertIn("confirmación", resp[0].lower())
        self.assertIn("Proyectos Durango", resp[0])
        resp_g = procesar_mensaje_web("TecnicoEstandar", "Guardar")
        self.assertIn("#099", resp_g[0])
        mock_crear.assert_called_once_with(
            ticket="TK-500",
            tecnico="TecnicoEstandar",
            ubicacion="Oficinas Centrales",
            actividad="Instalación Switch",
            estado="Terminado",
            departamento="Sistemas",
            evidencias=[],
            proyecto="Proyectos Durango",
        )

    @patch("app.db.crear_reporte", return_value="#100")
    def test_registro_rapido_retrocompatibilidad_sin_proyecto(self, mock_crear):
        resp = procesar_mensaje_web(
            "TecnicoEstandar",
            "/registro TK-100 | Redes | Nivel 10 | Cambio AP | Terminado"
        )
        self.assertIn("confirmación", resp[0].lower())
        self.assertIn("Mina San Dimas", resp[0])
        resp_g = procesar_mensaje_web("TecnicoEstandar", "Guardar")
        self.assertIn("#100", resp_g[0])
        mock_crear.assert_called_once_with(
            ticket="TK-100",
            tecnico="TecnicoEstandar",
            ubicacion="Nivel 10",
            actividad="Cambio AP",
            estado="Terminado",
            departamento="Redes",
            evidencias=[],
            proyecto="Mina San Dimas",
        )

    @patch("app.db.listar_proyectos")
    def test_comando_proyectos_lista(self, mock_proy):
        mock_proy.return_value = [
            {"nombre": "Mina San Dimas", "tipo": "Mina", "cliente": "First Majestic"},
            {"nombre": "Proyectos Durango", "tipo": "Durango", "cliente": "Qtek Computación"},
        ]
        resp = procesar_mensaje_web("TecnicoEstandar", "/proyectos")
        self.assertIn("Mina San Dimas", resp[0])
        self.assertIn("Proyectos Durango", resp[0])

    @patch("app.db.agregar_proyecto", return_value=True)
    def test_comando_nuevo_proyecto_admin(self, mock_add):
        resp = procesar_mensaje_web(
            "Miguel Abraham Lopez Ortiz",
            "/nuevo_proyecto Planta Durango | Durango | Qtek S.A."
        )
        mock_add.assert_called_once_with(
            nombre="Planta Durango",
            tipo="Durango",
            cliente="Qtek S.A."
        )
        self.assertIn("guardado correctamente", resp[0].lower())

    def test_excel_soporte_columna_proyecto(self):
        from app.excel import generar_excel
        reportes = [
            {
                "id": 1,
                "numero": "#001",
                "ticket": "TK-01",
                "tecnico": "Juan",
                "departamento": "Sistemas",
                "ubicacion": "Oficinas Centrales",
                "actividad": "Instalación",
                "estado": "Terminado",
                "evidencias": [],
                "fecha": "2026-10-20",
                "hora": "10:00",
                "actualizado": "2026-10-20 10:00:00",
                "proyecto": "Proyectos Durango",
            }
        ]
        wb_bytes, _ = generar_excel(reportes, dt.date(2026, 10, 20), dt.date(2026, 10, 20), proyecto="Proyectos Durango")
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
        ws_rep = wb["Reportes"]
        self.assertEqual(ws_rep["L1"].value, "Proyecto")
        self.assertEqual(ws_rep["L2"].value, "Proyectos Durango")
        ws_resumen = wb["Resumen"]
        texto_caratula = " ".join(str(ws_resumen.cell(row=r, column=2).value) for r in range(1, 10))
        self.assertIn("Qtek Computación", texto_caratula)

    @patch("app.db.obtener_proyectos_de_tecnico", return_value=["Mina San Dimas"])
    def test_tecnico_con_un_solo_proyecto_asigna_automatico(self, mock_proy):
        estado = get_estado("TecnicoMina")
        resp = procesar_mensaje_web("TecnicoMina", "➕ Nuevo reporte")
        self.assertEqual(estado.borrador.get("proyecto"), "Mina San Dimas")
        self.assertEqual(estado.esperando, "ticket")
        self.assertIn("Mina San Dimas", resp[0])
        self.assertIn("paso 2", resp[0].lower())

    @patch("app.db.obtener_proyectos_de_tecnico", return_value=["Proyectos Durango"])
    @patch("app.db.listar_nombres_proyectos", return_value=["Mina San Dimas", "Proyectos Durango"])
    def test_tecnico_no_puede_seleccionar_proyecto_no_asignado(self, mock_nombres, mock_proy):
        estado = get_estado("TecnicoDurango")
        estado.esperando = "proyecto"
        resp = procesar_mensaje_web("TecnicoDurango", "Mina San Dimas")
        self.assertIn("no tienes asignado", resp[0].lower())
        self.assertEqual(estado.esperando, "proyecto")

    @patch("app.db.asignar_proyecto_tecnico", return_value=True)
    def test_admin_asigna_proyecto_a_tecnico(self, mock_asig):
        resp = procesar_mensaje_web("Miguel Abraham Lopez Ortiz", "/asignar_proyecto Juan Pérez | Mina San Dimas")
        mock_asig.assert_called_once_with("Juan Pérez", "Mina San Dimas")
        self.assertIn("asignó el proyecto", resp[0].lower())

    @patch("app.db.listar_tecnicos_con_info")
    def test_admin_lista_tecnicos(self, mock_lista):
        mock_lista.return_value = [
            {"nombre": "Miguel Abraham Lopez Ortiz", "cargo": "Admin", "proyecto": "Ambos", "es_admin": True, "activo": True},
            {"nombre": "Juan Pérez", "cargo": "Técnico", "proyecto": "Mina San Dimas", "es_admin": False, "activo": True},
        ]
        resp = procesar_mensaje_web("Miguel Abraham Lopez Ortiz", "/tecnicos")
        mock_lista.assert_called_once()
        self.assertIn("Juan Pérez", resp[0])
        self.assertIn("Mina San Dimas", resp[0])

    @patch("app.db.listar_tecnicos", return_value=["Juan Pérez", "Carlos Ruiz"])
    @patch("app.db.listar_nombres_proyectos", return_value=["Mina San Dimas", "Proyectos Durango"])
    @patch("app.db.asignar_proyecto_tecnico", return_value=True)
    def test_admin_flujo_interactivo_asignar_proyecto(self, mock_asig, mock_proy, mock_tec):
        estado = get_estado("Miguel Abraham Lopez Ortiz")
        resp1 = procesar_mensaje_web("Miguel Abraham Lopez Ortiz", "👥 Asignar Proyecto")
        self.assertEqual(estado.esperando, "admin_sel_tecnico_asignar")
        self.assertIn("Juan Pérez", resp1[0])

        resp2 = procesar_mensaje_web("Miguel Abraham Lopez Ortiz", "Juan Pérez")
        self.assertEqual(estado.esperando, "admin_sel_proyecto_asignar")
        self.assertIn("Mina San Dimas", resp2[0])

        resp3 = procesar_mensaje_web("Miguel Abraham Lopez Ortiz", "Proyectos Durango")
        self.assertIsNone(estado.esperando)
        mock_asig.assert_called_once_with("Juan Pérez", "Proyectos Durango")
        self.assertIn("asignó el proyecto", resp3[0].lower())


if __name__ == "__main__":
    unittest.main()

