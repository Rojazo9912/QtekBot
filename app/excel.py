"""
Genera el reporte de actividades en Excel (.xlsx) a partir de los reportes
de la base de datos (ver db.reportes_en_periodo). Se arma en memoria cada vez
que el admin lo pide; no se guarda en el servidor.

Hojas:
- Resumen: datos del contrato, periodo y totales por estado, departamento y ubicación.
- Reportes: una fila por reporte, con filtros y las fotos en miniatura
  dentro de la celda "Evidencias".
- Evidencias: una fila por foto, con link.
"""
import datetime as dt
import io
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from typing import Optional

import httpx
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as PILImage

from app import storage
from app.config import CATALOGO_ESTADO_REPORTE, CONTRATO_INFO, ZONA_HORARIA

_AZUL = "1F3864"
_GRIS = "F2F2F2"
_COLOR_ESTADO = {
    "Terminado": "E2EFDA",
    "Pendiente": "FFF2CC",
    "No solucionado": "FCE4D6",
}

_FUENTE_TITULO = Font(bold=True, size=14, color=_AZUL)
_FUENTE_SECCION = Font(bold=True, size=11, color=_AZUL)
_FUENTE_ENCABEZADO = Font(bold=True, color="FFFFFF")
_RELLENO_ENCABEZADO = PatternFill("solid", fgColor=_AZUL)
_RELLENO_GRIS = PatternFill("solid", fgColor=_GRIS)
_BORDE = Border(bottom=Side(style="thin", color="D9D9D9"))
_ARRIBA = Alignment(vertical="top", wrap_text=True)

# Miniaturas de evidencia dentro de la hoja Reportes.
_MINI_ALTO_PX = 120
_MINI_ANCHO_MAX_PX = 200
_MINI_MARGEN_PX = 4
_MINI_MAX_POR_REPORTE = 6


def _descargar(url: str) -> Optional[bytes]:
    """Primero por Supabase Storage (funciona aunque el bucket sea privado);
    si falla, por el link público."""
    try:
        return storage.descargar_evidencia(url)
    except Exception as e:
        print(f"[excel] Storage no pudo descargar {url}: {e}")
    try:
        res = httpx.get(url, timeout=15, follow_redirects=True)
        res.raise_for_status()
        return res.content
    except Exception as e:
        print(f"[excel] no se pudo descargar la evidencia {url}: {e}")
        return None


def _miniatura(contenido: bytes) -> Optional[tuple[bytes, int, int]]:
    """Reduce la foto a _MINI_ALTO_PX de alto (JPEG). None si no es imagen
    (p. ej. un PDF): esa evidencia queda solo como link en la hoja Evidencias."""
    try:
        img = PILImage.open(io.BytesIO(contenido))
        img = img.convert("RGB")
    except Exception:
        return None
    img.thumbnail((_MINI_ANCHO_MAX_PX, _MINI_ALTO_PX))
    salida = io.BytesIO()
    img.save(salida, format="JPEG", quality=80)
    return salida.getvalue(), img.width, img.height


def _miniaturas_por_url(reportes: list[dict]) -> dict[str, tuple[bytes, int, int]]:
    urls = list(dict.fromkeys(
        url for r in reportes for url in r["evidencias"][:_MINI_MAX_POR_REPORTE]
    ))
    if not urls:
        return {}
    with ThreadPoolExecutor(max_workers=8) as pool:
        descargas = list(pool.map(_descargar, urls))
    miniaturas = {}
    for url, contenido in zip(urls, descargas):
        mini = _miniatura(contenido) if contenido else None
        if mini:
            miniaturas[url] = mini
    return miniaturas


def _encabezados(ws, fila: int, titulos: list[str], col_inicio: int = 1) -> None:
    for i, titulo in enumerate(titulos):
        c = ws.cell(row=fila, column=col_inicio + i, value=titulo)
        c.font = _FUENTE_ENCABEZADO
        c.fill = _RELLENO_ENCABEZADO
        c.alignment = Alignment(vertical="center", wrap_text=True)


def _anchos(ws, anchos: list[int]) -> None:
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def _tabla_conteo(ws, fila: int, titulo: str, encabezados: list[str], filas: list[list]) -> int:
    """Escribe una tabla pequeña con título; regresa la siguiente fila libre."""
    ws.cell(row=fila, column=1, value=titulo).font = _FUENTE_SECCION
    _encabezados(ws, fila + 1, encabezados)
    for i, valores in enumerate(filas):
        for j, v in enumerate(valores, start=1):
            c = ws.cell(row=fila + 2 + i, column=j, value=v)
            c.border = _BORDE
    if not filas:
        ws.cell(row=fila + 2, column=1, value="Sin reportes en el periodo").font = Font(italic=True, color="808080")
        return fila + 4
    return fila + 3 + len(filas)


def _hoja_resumen(ws, reportes: list[dict], desde: dt.date, hasta: dt.date) -> None:
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    _anchos(ws, [34, 16, 14, 14, 16])

    ws["A1"] = "Reporte de Actividades de Campo TI"
    ws["A1"].font = _FUENTE_TITULO

    generado = dt.datetime.now(ZONA_HORARIA).strftime("%Y-%m-%d %H:%M")
    datos = [
        ("Contratista", CONTRATO_INFO["contratista"]),
        ("Contrato marco No.", CONTRATO_INFO["contrato_marco_no"]),
        ("Ubicación de los servicios", CONTRATO_INFO["ubicacion_servicios"]),
        ("Área de servicio", CONTRATO_INFO["area_servicio"]),
        ("Periodo", f"{desde.isoformat()} a {hasta.isoformat()}"),
        ("Generado", generado),
        ("Total de reportes", len(reportes)),
    ]
    for i, (etiqueta, valor) in enumerate(datos, start=3):
        ws.cell(row=i, column=1, value=etiqueta).font = Font(bold=True)
        ws.cell(row=i, column=1).fill = _RELLENO_GRIS
        ws.cell(row=i, column=2, value=valor).alignment = Alignment(horizontal="left")

    fila = 3 + len(datos) + 1

    por_estado = Counter(r["estado"] for r in reportes)
    fila = _tabla_conteo(
        ws, fila, "Por estado", ["Estado", "Reportes"],
        [[e, por_estado.get(e, 0)] for e in CATALOGO_ESTADO_REPORTE] if reportes else [],
    )

    fila = _tabla_conteo(
        ws, fila, "Por departamento", ["Departamento", "Total", *CATALOGO_ESTADO_REPORTE],
        [
            ["Infraestructura", len(reportes)]
            + [sum(1 for r in reportes if r["estado"] == e) for e in CATALOGO_ESTADO_REPORTE]
        ] if reportes else [],
    )

    por_ubicacion = Counter(r["ubicacion"] for r in reportes)
    _tabla_conteo(
        ws, fila, "Por ubicación", ["Ubicación", "Reportes"],
        [[u, n] for u, n in por_ubicacion.most_common()],
    )


def _insertar_miniaturas(ws, fila: int, col: int, minis: list[tuple[bytes, int, int]]) -> int:
    """Pone las fotos una junto a otra dentro de la celda (fila, col).
    Se anclan a la celda, así que se ocultan con la fila al filtrar.
    Regresa el ancho total usado en píxeles."""
    x = _MINI_MARGEN_PX
    for datos, ancho, alto in minis:
        # Un BytesIO nuevo por imagen: openpyxl lo cierra al guardar, y la
        # misma foto puede aparecer en más de un reporte.
        img = XLImage(io.BytesIO(datos))
        img.width, img.height = ancho, alto
        img.anchor = TwoCellAnchor(
            editAs="twoCell",
            _from=AnchorMarker(col=col - 1, colOff=pixels_to_EMU(x),
                               row=fila - 1, rowOff=pixels_to_EMU(_MINI_MARGEN_PX)),
            to=AnchorMarker(col=col - 1, colOff=pixels_to_EMU(x + ancho),
                            row=fila - 1, rowOff=pixels_to_EMU(_MINI_MARGEN_PX + alto)),
        )
        ws.add_image(img)
        x += ancho + _MINI_MARGEN_PX
    return x


def _hoja_reportes(ws, reportes: list[dict]) -> None:
    titulos = [
        "Número", "Ticket", "Departamento", "Ubicación", "Actividad", "Estado",
        "Fotos", "Fecha", "Hora", "Última actualización", "Evidencias",
    ]
    _encabezados(ws, 1, titulos)
    _anchos(ws, [10, 16, 28, 14, 60, 16, 8, 12, 10, 20, 30])
    ws.freeze_panes = "A2"
    col_evidencias = len(titulos)
    miniaturas = _miniaturas_por_url(reportes)
    ancho_max_px = 0

    for i, r in enumerate(reportes, start=2):
        valores = [
            r["numero"], r["ticket"], "Infraestructura", r["ubicacion"], r["actividad"],
            r["estado"], len(r["evidencias"]), r["fecha"], r["hora"], r["actualizado"],
        ]
        for j, v in enumerate(valores, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = _ARRIBA
            c.border = _BORDE
        color = _COLOR_ESTADO.get(r["estado"])
        if color:
            ws.cell(row=i, column=6).fill = PatternFill("solid", fgColor=color)

        minis = [miniaturas[u] for u in r["evidencias"][:_MINI_MAX_POR_REPORTE] if u in miniaturas]
        if minis:
            ancho_px = _insertar_miniaturas(ws, i, col_evidencias, minis)
            ancho_max_px = max(ancho_max_px, ancho_px)
            # Puntos = píxeles * 3/4
            ws.row_dimensions[i].height = (_MINI_ALTO_PX + 2 * _MINI_MARGEN_PX) * 0.75
        elif r["evidencias"]:
            ws.cell(row=i, column=col_evidencias, value="Ver hoja Evidencias").alignment = _ARRIBA

    if ancho_max_px:
        # Ancho de columna en caracteres ≈ (píxeles - 5) / 7
        letra = get_column_letter(col_evidencias)
        ws.column_dimensions[letra].width = max(30, (ancho_max_px - 5) / 7 + 1)

    ultima = max(1, len(reportes)) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(titulos))}{ultima}"


def _hoja_evidencias(ws, reportes: list[dict]) -> None:
    _encabezados(ws, 1, ["Número", "Ticket", "Foto", "Link"])
    _anchos(ws, [10, 16, 8, 90])
    ws.freeze_panes = "A2"

    fila = 2
    for r in reportes:
        for n, url in enumerate(r["evidencias"], start=1):
            ws.cell(row=fila, column=1, value=r["numero"])
            ws.cell(row=fila, column=2, value=r["ticket"])
            ws.cell(row=fila, column=3, value=n)
            link = ws.cell(row=fila, column=4, value=url)
            link.hyperlink = url
            link.style = "Hyperlink"
            fila += 1
    if fila == 2:
        ws.cell(row=2, column=1, value="Sin fotos en el periodo").font = Font(italic=True, color="808080")


def generar_excel(reportes: list[dict], desde: dt.date, hasta: dt.date) -> tuple[bytes, str]:
    """Regresa (contenido del .xlsx, nombre de archivo sugerido)."""
    wb = Workbook()
    _hoja_resumen(wb.active, reportes, desde, hasta)
    _hoja_reportes(wb.create_sheet("Reportes"), reportes)
    _hoja_evidencias(wb.create_sheet("Evidencias"), reportes)

    buffer = io.BytesIO()
    wb.save(buffer)
    nombre = f"Reporte_Campo_{desde.isoformat()}_a_{hasta.isoformat()}.xlsx"
    return buffer.getvalue(), nombre
